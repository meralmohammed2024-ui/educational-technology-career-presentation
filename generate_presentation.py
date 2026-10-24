from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

OUTPUT_FILE = "Educational_Technology_Career_Opportunities_Premium_V2.pptx"

NAVY = RGBColor(16, 24, 40)
BLUE = RGBColor(59, 130, 246)
TEAL = RGBColor(20, 184, 166)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(243, 244, 246)
GRAY = RGBColor(107, 114, 128)
DARK = RGBColor(31, 41, 55)
PANEL = RGBColor(19, 33, 57)
MUTED = RGBColor(212, 218, 228)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide, color=LIGHT):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, text, x=0.7, y=0.35, w=10.0, h=0.6, size=28, color=NAVY, bold=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, text, x, y, w, h, size=18, color=DARK, bold=False, italic=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_bullets(slide, items, x, y, w, h, size=18, color=DARK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.text = item
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        p.runs[0].font.name = "Aptos"
        p.runs[0].font.size = Pt(size)
        p.runs[0].font.color.rgb = color


def add_card(slide, x, y, w, h, fill, title, subtitle=None, title_size=16, title_color=WHITE, sub_color=LIGHT):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    shape.line.width = Pt(1)

    title_box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.12), Inches(w - 0.36), Inches(0.35))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    p.runs[0].font.name = "Aptos"
    p.runs[0].font.size = Pt(title_size)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = title_color

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.48), Inches(w - 0.38), Inches(h - 0.56))
        q = sub_box.text_frame.paragraphs[0]
        q.text = subtitle
        q.alignment = PP_ALIGN.LEFT
        q.runs[0].font.name = "Aptos"
        q.runs[0].font.size = Pt(10)
        q.runs[0].font.color.rgb = sub_color


def add_small_label(slide, x, y, w, h, text, fill=BLUE, color=WHITE):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    rect.fill.solid(); rect.fill.fore_color.rgb = fill; rect.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(x + 0.1), Inches(y + 0.04), Inches(w - 0.2), Inches(h - 0.08))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.name = "Aptos"
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = color


def add_footer(slide, text):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.12), Inches(13.333), Inches(0.38))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(7.2), Inches(12.0), Inches(0.2))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.runs[0].font.name = "Aptos"
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = WHITE


# Slide 1 - Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)

for x, y, w, h, col in [
    (0.5, 1.0, 2.7, 1.3, BLUE),
    (3.5, 1.6, 2.5, 1.4, TEAL),
    (6.2, 1.1, 1.8, 1.2, BLUE),
    (8.7, 1.7, 2.7, 1.3, TEAL),
    (1.2, 3.5, 3.3, 1.8, BLUE),
    (4.9, 3.9, 2.7, 1.2, TEAL),
    (7.9, 3.6, 3.5, 1.6, BLUE),
]:
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    rect.fill.solid(); rect.fill.fore_color.rgb = col
    rect.line.fill.background()

add_title(slide, "Career Opportunities in Educational Technology", x=0.8, y=1.0, w=8.3, size=29, color=WHITE)
add_title(slide, "Exploring Professional Roles in the Digital Learning Era", x=0.9, y=2.0, w=8.5, size=18, color=LIGHT, bold=False)

pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.55), Inches(2.5), Inches(0.4))
pill.fill.solid(); pill.fill.fore_color.rgb = TEAL; pill.line.fill.background()
box = slide.shapes.add_textbox(Inches(1.05), Inches(2.64), Inches(2.1), Inches(0.2))
p = box.text_frame.paragraphs[0]
p.text = "University learners"
p.alignment = PP_ALIGN.CENTER
p.runs[0].font.name = "Aptos"; p.runs[0].font.size = Pt(11); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = WHITE

panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.0), Inches(2.8), Inches(3.4), Inches(2.0))
panel.fill.solid(); panel.fill.fore_color.rgb = PANEL; panel.line.color.rgb = PANEL
add_text(slide, "Educational Technology", 9.25, 3.2, 2.8, 0.35, size=18, color=WHITE, bold=True)
add_text(slide, "Instructional design • Digital learning • Data • AI • Inclusion", 9.25, 3.8, 2.8, 0.65, size=10, color=LIGHT)

add_footer(slide, "Academic overview • Career pathways • Professional development")

# Slide 2 - Intro
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, "Introduction to Educational Technology", size=27, color=NAVY)
line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.98), Inches(2.4), Inches(0.06))
line.fill.solid(); line.fill.fore_color.rgb = BLUE; line.line.fill.background()

add_text(slide,
    "Educational Technology is the systematic design, development, implementation, management, and evaluation of learning processes and resources.",
    0.8, 1.3, 5.6, 1.0, size=20, color=DARK)
add_text(slide,
    "It integrates pedagogy, technology, and evidence-based practice to improve teaching, learning, and assessment, extending beyond simply using digital devices in classrooms.",
    0.8, 2.2, 5.7, 1.3, size=18, color=DARK)

add_card(slide, 7.0, 1.35, 2.2, 1.25, BLUE, "Design", "Learning goals, tasks, and strategies")
add_card(slide, 9.5, 1.35, 2.5, 1.25, TEAL, "Development", "Resources and digital content")
add_card(slide, 7.0, 2.9, 2.2, 1.25, BLUE, "Implementation", "Delivery and support")
add_card(slide, 9.5, 2.9, 2.5, 1.25, TEAL, "Evaluation", "Evidence and improvement")

quote = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.2), Inches(11.4), Inches(0.9))
quote.fill.solid(); quote.fill.fore_color.rgb = WHITE; quote.line.color.rgb = RGBColor(220, 223, 227)
add_text(slide, "Educational technology is not only about tools; it is about creating meaningful, effective, and equitable learning experiences.", 1.2, 5.45, 10.6, 0.5, size=15, color=DARK, italic=True)
add_footer(slide, "Field definition and scope")

# Slide 3 - Why it matters
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, "Why Is Educational Technology Important?", size=27, color=WHITE)

cards = [
    ("Instructional design", "Clear objectives and learning strategies", 0.8, 1.4, 2.2, 1.2, BLUE),
    ("Accessibility", "Inclusive learning support", 3.25, 1.4, 2.3, 1.2, TEAL),
    ("Personalized learning", "Adaptive pathways and pacing", 5.8, 1.4, 2.5, 1.2, BLUE),
    ("Assessment", "Evidence and feedback", 8.55, 1.4, 2.4, 1.2, TEAL),
    ("Professional training", "Workforce learning", 10.8, 1.4, 1.7, 1.2, BLUE),
    ("Content development", "Interactive and multimedia resources", 1.7, 3.1, 2.7, 1.4, TEAL),
    ("Data literacy", "Data-informed decisions", 4.75, 3.1, 2.5, 1.4, BLUE),
    ("Collaboration", "Digital teamwork and communication", 7.55, 3.1, 2.4, 1.4, TEAL),
    ("Innovation", "Future-ready learning design", 10.1, 3.1, 2.3, 1.4, BLUE),
]
for title, sub, x, y, w, h, col in cards:
    add_card(slide, x, y, w, h, col, title, sub, title_size=15)

summary = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(5.25), Inches(10.8), Inches(0.75))
summary.fill.solid(); summary.fill.fore_color.rgb = RGBColor(17, 41, 69); summary.line.fill.background()
add_text(slide, "Educational technology supports learning quality, inclusion, efficiency, evidence, and innovation across formal and informal settings.", 1.5, 5.42, 10.3, 0.3, size=14, color=WHITE)
add_footer(slide, "Why it matters")

# Slide 4 - Landscape
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, "The Educational Technology Career Landscape", size=27, color=NAVY)

entries = [
    ("Instructional Design", "Goals, strategies, assessments", 0.8, 1.2, 2.5, 1.0, BLUE),
    ("Digital Learning", "Course and content design", 3.6, 1.2, 2.4, 1.0, TEAL),
    ("LMS & Support", "Systems and reporting", 6.3, 1.2, 2.8, 1.0, BLUE),
    ("Online Education", "Virtual delivery and blended learning", 9.5, 1.2, 2.6, 1.0, TEAL),
    ("Learning Analytics", "Data patterns and improvement", 0.8, 2.7, 2.8, 1.0, TEAL),
    ("Multimedia & Simulations", "Video, animation, and interactivity", 3.9, 2.7, 2.7, 1.0, BLUE),
    ("Research & Evaluation", "Usability and outcomes", 6.9, 2.7, 2.9, 1.0, TEAL),
    ("Management & Consulting", "Strategy and implementation", 10.0, 2.7, 2.2, 1.0, BLUE),
]
for title, sub, x, y, w, h, col in entries:
    add_card(slide, x, y, w, h, col, title, sub, title_size=15)

center = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.9), Inches(4.2), Inches(5.5), Inches(1.8))
center.fill.solid(); center.fill.fore_color.rgb = WHITE; center.line.color.rgb = RGBColor(208, 217, 233)
add_text(slide, "Educational technology professionals connect learning science, digital tools, evidence, and user experience", 4.15, 4.45, 5.1, 1.0, size=18, color=NAVY, bold=True)
for x, label in [(2.2, "Learner needs"), (5.7, "Design"), (8.5, "Tools"), (10.75, "Impact")]:
    add_small_label(slide, x, 6.15, 1.25, 0.35, label, fill=BLUE)
add_footer(slide, "Major professional categories")

# Slide 5 - Instructional Designer
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, "Instructional Designer", size=28, color=WHITE)

add_bullets(slide, [
    "Analyze learning needs and define measurable objectives.",
    "Design instructional strategies, activities, and assessments.",
    "Develop learning resources and align them to learning outcomes.",
    "Evaluate learning effectiveness and revise instruction.",
], 0.8, 1.3, 5.8, 2.9, size=18, color=WHITE)

right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.4), Inches(5.2), Inches(3.8))
right.fill.solid(); right.fill.fore_color.rgb = RGBColor(22, 36, 58); right.line.color.rgb = RGBColor(22, 36, 58)
steps = [
    ("Analyze", "Needs, audience, context"),
    ("Design", "Objectives and strategy"),
    ("Develop", "Content and digital resources"),
    ("Implement", "Delivery and facilitation"),
    ("Evaluate", "Feedback and evidence"),
]
start_y = 1.7
for idx, (title, subtitle) in enumerate(steps):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.45), Inches(start_y), Inches(4.3), Inches(0.5))
    rect.fill.solid(); rect.fill.fore_color.rgb = BLUE if idx % 2 == 0 else TEAL; rect.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(7.7), Inches(start_y + 0.06), Inches(3.9), Inches(0.27))
    p = tb.text_frame.paragraphs[0]
    p.text = f"{title}: {subtitle}"
    p.runs[0].font.name = "Aptos"; p.runs[0].font.size = Pt(11); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = WHITE
    start_y += 0.62

call = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.7), Inches(11.4), Inches(0.75))
call.fill.solid(); call.fill.fore_color.rgb = WHITE; call.line.fill.background()
add_text(slide, "Common models include ADDIE and backward design, which help align goals, activities, and assessment.", 1.2, 5.9, 10.8, 0.24, size=14, color=NAVY)
add_footer(slide, "Instructional Designer")

# Slide 6 - LXD
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, "Learning Experience Designer (LXD)", size=27, color=NAVY)

left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(4.6), Inches(4.5))
left.fill.solid(); left.fill.fore_color.rgb = WHITE; left.line.color.rgb = RGBColor(219, 227, 239)
add_text(slide, "Learner-centered design", 1.1, 1.5, 3.7, 0.4, size=20, color=NAVY, bold=True)
add_bullets(slide, [
    "Learner needs and motivations",
    "Usability and accessibility",
    "Interaction and engagement",
    "Clarity of content and tasks",
    "Instructional goals and outcomes",
], 1.1, 2.0, 3.7, 2.4, size=16, color=DARK)

middle = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(5.55), Inches(1.9), Inches(2.15), Inches(1.8))
middle.fill.solid(); middle.fill.fore_color.rgb = BLUE; middle.line.fill.background()
add_text(slide, "UX +\npedagogy", 5.7, 2.3, 1.7, 0.75, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.1), Inches(1.2), Inches(4.4), Inches(4.5))
right.fill.solid(); right.fill.fore_color.rgb = TEAL; right.line.color.rgb = TEAL
add_text(slide, "Key design principles", 8.5, 1.5, 3.3, 0.4, size=20, color=WHITE, bold=True)
add_bullets(slide, [
    "Accessibility and inclusion",
    "Clear structure and flow",
    "Feedback and reflection",
    "Motivation and relevance",
    "Evidence-informed iteration",
], 8.5, 2.0, 3.2, 2.2, size=16, color=WHITE)

summary = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(5.9), Inches(10.9), Inches(0.7))
summary.fill.solid(); summary.fill.fore_color.rgb = WHITE; summary.line.fill.background()
add_text(slide, "A strong learning experience balances learner needs, usability, evidence, and meaningful interaction.", 1.5, 6.09, 10.1, 0.25, size=14, color=NAVY, italic=True)
add_footer(slide, "Learning Experience Designer")

# Slide 7 - E-Learning Developer
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, "E-Learning Developer", size=28, color=WHITE)
add_bullets(slide, [
    "Build interactive digital courses and modules.",
    "Create quizzes, branching scenarios, and simulations.",
    "Develop engaging learning experiences aligned with outcomes.",
    "Use authoring tools such as Articulate Storyline 360, Rise 360, and H5P.",
], 0.8, 1.3, 5.9, 3.0, size=18, color=WHITE)

right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(1.4), Inches(5.1), Inches(3.8))
right.fill.solid(); right.fill.fore_color.rgb = RGBColor(20, 40, 70); right.line.color.rgb = RGBColor(20, 40, 70)
items = [
    ("Articulate", "Storyline 360"),
    ("Articulate", "Rise 360"),
    ("H5P", "Interactive content"),
    ("SCORM / LMS", "Packaging and delivery"),
    ("Accessibility", "Inclusive design"),
]
start_y = 1.8
for idx, (title, subtitle) in enumerate(items):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.45), Inches(start_y), Inches(4.45), Inches(0.5))
    rect.fill.solid(); rect.fill.fore_color.rgb = BLUE if idx % 2 == 0 else TEAL; rect.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(7.7), Inches(start_y + 0.06), Inches(3.9), Inches(0.28))
    p = tb.text_frame.paragraphs[0]
    p.text = f"{title}: {subtitle}"
    p.runs[0].font.name = "Aptos"; p.runs[0].font.size = Pt(11); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = WHITE
    start_y += 0.62

note = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.75), Inches(11.4), Inches(0.75))
note.fill.solid(); note.fill.fore_color.rgb = WHITE; note.line.fill.background()
add_text(slide, "E-learning developers combine instructional design, user experience, and technical implementation.", 1.2, 6.0, 10.8, 0.2, size=14, color=NAVY)
add_footer(slide, "E-Learning Developer")

# Slide 8 - Multimedia
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, "Educational Multimedia and Digital Content Specialist", size=25, color=NAVY)
cards = [
    ("Video", "Instructional videos and explainers", 0.8, 1.4, 2.3, 1.5, BLUE),
    ("Animation", "Motion graphics and visual explanations", 3.45, 1.4, 2.3, 1.5, TEAL),
    ("Graphics", "Illustrations and infographics", 6.1, 1.4, 2.3, 1.5, BLUE),
    ("Podcasts", "Audio learning and storytelling", 8.75, 1.4, 2.3, 1.5, TEAL),
    ("Interactive media", "Dynamic learning tools and simulations", 1.7, 3.45, 3.1, 1.7, TEAL),
    ("Presentations", "Narrative decks and communication", 5.15, 3.45, 3.1, 1.7, BLUE),
    ("Curriculum materials", "Teacher-friendly resources", 8.7, 3.45, 3.0, 1.7, TEAL),
]
for title, subtitle, x, y, w, h, col in cards:
    add_card(slide, x, y, w, h, col, title, subtitle, title_size=15)
note = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(5.8), Inches(10.9), Inches(0.7))
note.fill.solid(); note.fill.fore_color.rgb = WHITE; note.line.fill.background()
add_text(slide, "These roles require strong storytelling, media production, communication, and instructional awareness.", 1.5, 6.02, 10.2, 0.25, size=14, color=NAVY)
add_footer(slide, "Multimedia and digital content")

# Slide 9 - LMS
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, "Learning Management System (LMS) Specialist", size=26, color=WHITE)
add_bullets(slide, [
    "Manage course shells, upload content, and organize learning spaces.",
    "Configure user accounts, permissions, and enrollment processes.",
    "Support assessment workflows, grading, and reporting.",
    "Troubleshoot platform issues and support instructors and learners.",
    "Evaluate system performance and improve digital learning operations.",
], 0.8, 1.2, 5.9, 3.2, size=18, color=WHITE)

right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.4), Inches(5.1), Inches(3.7))
right.fill.solid(); right.fill.fore_color.rgb = RGBColor(20, 40, 70); right.line.color.rgb = RGBColor(20, 40, 70)
items = [
    ("Moodle", "Open-source LMS"),
    ("Canvas", "Learning platform"),
    ("Blackboard", "Institutional LMS"),
    ("Reporting", "Analytics and tracking"),
    ("Support", "Troubleshooting and facilitation"),
]
start_y = 1.8
for idx, (title, subtitle) in enumerate(items):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.35), Inches(start_y), Inches(4.4), Inches(0.5))
    rect.fill.solid(); rect.fill.fore_color.rgb = BLUE if idx % 2 == 0 else TEAL; rect.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(7.6), Inches(start_y + 0.06), Inches(3.9), Inches(0.27))
    p = tb.text_frame.paragraphs[0]
    p.text = f"{title}: {subtitle}"
    p.runs[0].font.name = "Aptos"; p.runs[0].font.size = Pt(11); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = WHITE
    start_y += 0.62
add_footer(slide, "LMS and platform support")

# Slide 10 - Online learning
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, "Online Learning and Virtual Classroom Specialist", size=25, color=NAVY)
items = [
    ("Course design", "Structure, pacing, and accessibility", 0.8, 1.4, 5.2, 1.5, BLUE),
    ("Virtual facilitation", "Live sessions and moderation", 6.2, 1.4, 5.2, 1.5, TEAL),
    ("Blended learning", "Coordinating online and in-person experiences", 0.8, 3.2, 5.2, 1.5, TEAL),
    ("Community building", "Discussion, interaction, and learner connection", 6.2, 3.2, 5.2, 1.5, BLUE),
]
for title, subtitle, x, y, w, h, col in items:
    add_card(slide, x, y, w, h, col, title, subtitle, title_size=15)
note = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(5.9), Inches(10.9), Inches(0.6))
note.fill.solid(); note.fill.fore_color.rgb = WHITE; note.line.fill.background()
add_text(slide, "This role emphasizes pedagogy, communication, digital collaboration, and learner support in online settings.", 1.5, 6.08, 10.2, 0.25, size=14, color=NAVY)
add_footer(slide, "Online and virtual learning")

# Slide 11 - AI
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, "AI in Education Specialist", size=28, color=WHITE)
add_bullets(slide, [
    "Use AI to support tutoring, adaptive learning, and individualized feedback.",
    "Create or evaluate AI-assisted content generation and course support tools.",
    "Design ethical and accountable AI integrations with clear learner safeguards.",
    "Maintain human oversight, privacy protections, accessibility, and transparency.",
], 0.8, 1.3, 5.9, 3.1, size=18, color=WHITE)
right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.4), Inches(5.1), Inches(3.7))
right.fill.solid(); right.fill.fore_color.rgb = RGBColor(24, 48, 70); right.line.color.rgb = RGBColor(24, 48, 70)
items = [
    ("AI tutoring", "Adaptive learning support"),
    ("Automated feedback", "Rapid formative assessment"),
    ("Content assistance", "Drafting and ideation"),
    ("Ethics & privacy", "Human oversight and safety"),
]
start_y = 1.8
for idx, (title, subtitle) in enumerate(items):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.35), Inches(start_y), Inches(4.4), Inches(0.55))
    rect.fill.solid(); rect.fill.fore_color.rgb = BLUE if idx % 2 == 0 else TEAL; rect.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(7.6), Inches(start_y + 0.06), Inches(3.9), Inches(0.28))
    p = tb.text_frame.paragraphs[0]
    p.text = f"{title}: {subtitle}"
    p.runs[0].font.name = "Aptos"; p.runs[0].font.size = Pt(11); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = WHITE
    start_y += 0.65
banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.8), Inches(11.3), Inches(0.7))
banner.fill.solid(); banner.fill.fore_color.rgb = WHITE; banner.line.fill.background()
add_text(slide, "Human oversight, privacy, accessibility, and responsible use remain foundational in AI-supported education.", 1.2, 6.02, 10.7, 0.25, size=14, color=NAVY)
add_footer(slide, "AI in education")

# Slide 12 - Learning analytics
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, "Learning Analytics and Educational Data Specialist", size=25, color=NAVY)
add_bullets(slide, [
    "Analyze participation, engagement, completion, and assessment patterns.",
    "Use dashboards to understand learner behavior and course performance.",
    "Support instructional improvement through evidence-informed decisions.",
    "Translate learning data into actionable recommendations without over-claiming causation.",
], 0.8, 1.3, 5.8, 2.8, size=18, color=DARK)
panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(1.3), Inches(5.3), Inches(3.8))
panel.fill.solid(); panel.fill.fore_color.rgb = WHITE; panel.line.color.rgb = RGBColor(220, 224, 230)
for idx, h in enumerate([0.5, 1.1, 1.5, 2.2, 2.8, 3.1]):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.5 + idx * 0.7), Inches(4.5 - h), Inches(0.45), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE if idx % 2 == 0 else TEAL; bar.line.fill.background()
line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.35), Inches(4.6), Inches(4.5), Inches(0.05))
line.fill.solid(); line.fill.fore_color.rgb = GRAY; line.line.fill.background()
for label, x in [("Engage", 7.52), ("Complete", 8.2), ("Quiz", 8.9), ("Feedback", 9.6), ("Progress", 10.35), ("Support", 11.05)]:
    box = slide.shapes.add_textbox(Inches(x), Inches(4.8), Inches(0.58), Inches(0.2))
    p = box.text_frame.paragraphs[0]
    p.text = label
    p.runs[0].font.name = "Aptos"; p.runs[0].font.size = Pt(8); p.runs[0].font.color.rgb = DARK
note = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.9), Inches(11.5), Inches(0.6))
note.fill.solid(); note.fill.fore_color.rgb = WHITE; note.line.fill.background()
add_text(slide, "Examples of dashboard indicators include engagement, completion, performance trends, and learner support needs.", 1.3, 6.08, 10.8, 0.25, size=14, color=NAVY)
add_footer(slide, "Learning analytics")

# Slide 13 - Researcher and Evaluator
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, "Educational Technology Researcher and Evaluator", size=25, color=WHITE)
add_bullets(slide, [
    "Design studies and evaluate educational interventions and technologies.",
    "Investigate usability, learner experience, and learning outcomes.",
    "Use qualitative and quantitative methods to assess effectiveness.",
    "Produce evidence that informs course design, policy, and practice.",
], 0.8, 1.3, 5.8, 2.8, size=18, color=WHITE)
right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(1.3), Inches(5.2), Inches(3.8))
right.fill.solid(); right.fill.fore_color.rgb = RGBColor(24, 48, 70); right.line.color.rgb = RGBColor(24, 48, 70)
steps = [
    ("Problem", "Identify a learning challenge"),
    ("Study design", "Methods and measures"),
    ("Data", "Collect evidence"),
    ("Analysis", "Interpret results"),
    ("Improve", "Inform design decisions"),
]
start_y = 1.7
for idx, (title, subtitle) in enumerate(steps):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.35), Inches(start_y), Inches(4.5), Inches(0.55))
    rect.fill.solid(); rect.fill.fore_color.rgb = BLUE if idx % 2 == 0 else TEAL; rect.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(7.6), Inches(start_y + 0.06), Inches(4.0), Inches(0.28))
    p = tb.text_frame.paragraphs[0]
    p.text = f"{title}: {subtitle}"
    p.runs[0].font.name = "Aptos"; p.runs[0].font.size = Pt(11); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = WHITE
    start_y += 0.65
add_footer(slide, "Research and evaluation")

# Slide 14 - Emerging paths
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, "Emerging and Specialized Career Paths", size=27, color=NAVY)
entries = [
    ("VR / AR learning", "Immersive environments and simulations", BLUE),
    ("Educational game design", "Gamification and engagement", TEAL),
    ("Digital accessibility", "Inclusive design and assistive technology", BLUE),
    ("Learning technology consulting", "Strategy and implementation support", TEAL),
    ("Project management", "EdTech rollout and governance", BLUE),
]
positions = [(0.8, 1.3), (0.8, 2.8), (0.8, 4.2), (6.2, 1.3), (6.2, 2.8)]
for i, (title, subtitle, col) in enumerate(entries):
    x, y = positions[i]
    add_card(slide, x, y, 4.9, 1.2, col, title, subtitle, title_size=16)
label = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.2), Inches(5.5), Inches(4.8), Inches(1.0))
label.fill.solid(); label.fill.fore_color.rgb = WHITE; label.line.color.rgb = RGBColor(200, 200, 200)
add_text(slide, "Established roles and emerging specializations often overlap and evolve together.", 4.45, 5.8, 4.1, 0.4, size=15, color=NAVY)
add_footer(slide, "Emerging opportunities")

# Slide 15 - Skills
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, "Essential Professional Skills", size=28, color=WHITE)
skills = [
    ("Instructional design", 0.8, 1.3, 2.5, 1.15, BLUE),
    ("Digital content", 3.5, 1.3, 2.3, 1.15, TEAL),
    ("Platforms & software", 6.1, 1.3, 2.3, 1.15, BLUE),
    ("Data literacy", 8.8, 1.3, 2.25, 1.15, TEAL),
    ("Communication", 0.8, 2.9, 2.5, 1.15, TEAL),
    ("Project management", 3.55, 2.9, 2.5, 1.15, BLUE),
    ("Accessibility", 6.3, 2.9, 2.4, 1.15, TEAL),
    ("Adaptability", 9.0, 2.9, 2.1, 1.15, BLUE),
]
for title, x, y, w, h, col in skills:
    add_card(slide, x, y, w, h, col, title, "", title_size=15)
panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(5.1), Inches(10.8), Inches(0.85))
panel.fill.solid(); panel.fill.fore_color.rgb = WHITE; panel.line.fill.background()
add_text(slide, "Strong educational technology professionals combine pedagogy, digital competence, creativity, evidence, and ethical judgment.", 1.5, 5.38, 10.2, 0.3, size=15, color=NAVY)
add_footer(slide, "Core professional skills")

# Slide 16 - Work environments
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, "Where Can Educational Technology Professionals Work?", size=26, color=NAVY)
entries = [
    ("Universities and schools", "Teaching, course design, and support", BLUE),
    ("E-learning companies", "Product design and content development", TEAL),
    ("Corporate L&D", "Training and workforce development", BLUE),
    ("Government & agencies", "Policy and digital learning initiatives", TEAL),
    ("Research institutions", "Evaluation and learning studies", BLUE),
    ("Consulting teams", "Strategy and implementation support", TEAL),
    ("Freelance and remote work", "Independent digital learning services", BLUE),
]
positions = [(0.8, 1.3), (6.0, 1.3), (0.8, 2.9), (6.0, 2.9), (0.8, 4.5), (6.0, 4.5), (3.4, 5.9)]
for idx, (title, subtitle, col) in enumerate(entries):
    x, y = positions[idx]
    if title == "Freelance and remote work":
        add_card(slide, x, y, 5.9, 1.3, col, title, subtitle, title_size=15)
    else:
        add_card(slide, x, y, 4.8, 1.3, col, title, subtitle, title_size=15)
add_footer(slide, "Employment contexts")

# Slide 17 - Roadmap
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, "Career Roadmap", size=28, color=WHITE)
steps = [
    "1. Identify personal interests and strengths",
    "2. Learn instructional design fundamentals",
    "3. Develop proficiency in relevant digital tools",
    "4. Create a portfolio of educational projects",
    "5. Gain practical experience through internships, projects, or volunteering",
    "6. Demonstrate accessibility, evaluation, and ethical design skills",
    "7. Continue learning and specialize in a career path",
]
start_y = 1.2
for idx, step in enumerate(steps):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(start_y), Inches(5.9), Inches(0.52))
    rect.fill.solid(); rect.fill.fore_color.rgb = BLUE if idx % 2 == 0 else TEAL; rect.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(1.1), Inches(start_y + 0.08), Inches(5.4), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    p.text = step
    p.runs[0].font.name = "Aptos"; p.runs[0].font.size = Pt(14); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = WHITE
    start_y += 0.66

portfolio = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(1.4), Inches(4.9), Inches(4.6))
portfolio.fill.solid(); portfolio.fill.fore_color.rgb = RGBColor(20, 40, 70); portfolio.line.color.rgb = RGBColor(20, 40, 70)
add_text(slide, "Sample portfolio", 7.5, 1.8, 3.8, 0.28, size=18, color=WHITE, bold=True)
for i, item in enumerate([
    "Instructional design document",
    "Interactive learning module",
    "Multimedia resource",
    "Short evaluation report",
]):
    add_text(slide, f"• {item}", 7.55, 2.6 + i * 0.7, 4.0, 0.3, size=12, color=WHITE)
add_footer(slide, "Career roadmap")

# Slide 18 - Conclusion and references
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, "Conclusion and References", size=28, color=NAVY)
summary_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.1), Inches(7.1), Inches(1.7))
summary_box.fill.solid(); summary_box.fill.fore_color.rgb = WHITE; summary_box.line.color.rgb = RGBColor(220, 224, 231)
add_text(slide, "Educational technology is a diverse and evolving field that brings together pedagogy, creativity, technology, research, and social impact. Careers in this area require both intellectual rigor and practical innovation.", 1.1, 1.5, 6.4, 0.95, size=18, color=NAVY)
refs = (
    "Association for Educational Communications and Technology (AECT): https://www.aect.org/\n"
    "International Society for Technology in Education (ISTE): https://www.iste.org/\n"
    "EDUCAUSE: https://www.educause.edu/\n"
    "U.S. Department of Education, Office of EdTech: https://tech.ed.gov/\n"
    "UNESCO ICT in Education: https://en.unesco.org/themes/ict-education\n"
    "European Commission, Digital Education: https://education.ec.europa.eu/\n"
    "Open Educational Resources (OER Commons): https://www.oercommons.org/"
)
ref_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.7), Inches(3.0))
ref_box.text_frame.word_wrap = True
p = ref_box.text_frame.paragraphs[0]
p.text = refs
p.runs[0].font.name = "Aptos"; p.runs[0].font.size = Pt(12); p.runs[0].font.color.rgb = DARK
end_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.5))
end_bar.fill.solid(); end_bar.fill.fore_color.rgb = RGBColor(17, 41, 69); end_bar.line.fill.background()
add_text(slide, "Career opportunities in educational technology are broad, dynamic, and increasingly important in digital learning.", 1.1, 6.33, 11.0, 0.2, size=13, color=WHITE)
add_footer(slide, "Conclusion and references")

prs.save(OUTPUT_FILE)
print(f"Created premium version: {OUTPUT_FILE}")
